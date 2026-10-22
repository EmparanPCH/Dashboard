"""Crea un extracto de solo valores para publicar el dashboard sin cargar el Excel."""

import argparse
import base64
from pathlib import Path

from openpyxl import Workbook, load_workbook


NOMBRE_HOJA = "DUNE BESS CONSTR."
FILA_INICIO_EXTRACTO = 8
COL_ULTIMA = 158  # FB
RANGOS_REQUERIDOS = ((2, 3), (6, 8), (10, 12), (18, 19), (23, 89), (92, 158))


def preparar_datos(origen: Path, destino: Path) -> int:
    """Copia solo las celdas utilizadas por el dashboard, sin macros ni otras hojas."""
    libro_origen = load_workbook(
        origen,
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    try:
        if NOMBRE_HOJA not in libro_origen.sheetnames:
            hojas = ", ".join(libro_origen.sheetnames)
            raise ValueError(f'No se encontró la hoja "{NOMBRE_HOJA}". Hojas: {hojas}')

        hoja_origen = libro_origen[NOMBRE_HOJA]
        ultima_fila = hoja_origen.max_row or FILA_INICIO_EXTRACTO
        libro_extracto = Workbook()
        hoja_extracto = libro_extracto.active
        hoja_extracto.title = NOMBRE_HOJA

        for numero_fila, valores in enumerate(
            hoja_origen.iter_rows(
                min_row=FILA_INICIO_EXTRACTO,
                max_row=ultima_fila,
                min_col=1,
                max_col=COL_ULTIMA,
                values_only=True,
            ),
            start=FILA_INICIO_EXTRACTO,
        ):
            for inicio, fin in RANGOS_REQUERIDOS:
                for columna in range(inicio, fin + 1):
                    valor = valores[columna - 1]
                    if valor is not None:
                        hoja_extracto.cell(numero_fila, columna, valor)

        destino.parent.mkdir(parents=True, exist_ok=True)
        libro_extracto.save(destino)
        libro_extracto.close()
        return ultima_fila
    finally:
        libro_origen.close()


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Prepara los datos de construcción para el dashboard público."
    )
    parser.add_argument(
        "archivo_excel",
        type=Path,
        help="Ruta del archivo .xlsm o .xlsx de origen.",
    )
    parser.add_argument(
        "--salida",
        type=Path,
        default=Path(__file__).resolve().with_name("dashboard_data.xlsx"),
        help="Ruta de salida del extracto .xlsx.",
    )
    args = parser.parse_args()

    if not args.archivo_excel.is_file():
        parser.error(f"No existe el archivo de origen: {args.archivo_excel}")
    if args.archivo_excel.resolve() == args.salida.resolve():
        parser.error("El archivo de salida no puede reemplazar el original.")

    filas = preparar_datos(args.archivo_excel, args.salida)
    ruta_codificada = Path(f"{args.salida}.b64")
    ruta_codificada.write_text(
        base64.b64encode(args.salida.read_bytes()).decode("ascii"),
        encoding="ascii",
    )
    tamaño_mb = args.salida.stat().st_size / (1024 * 1024)
    print(f"Extracto creado: {args.salida} ({tamaño_mb:.2f} MB; hasta fila {filas}).")
    print(f"Copia preparada para Streamlit Cloud: {ruta_codificada}.")
    print("El extracto conserva solo los valores necesarios de la hoja del dashboard.")


if __name__ == "__main__":
    main()

