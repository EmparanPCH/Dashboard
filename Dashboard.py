from datetime import date, datetime
from io import BytesIO
from html import escape
from collections import Counter
import math
import re
import unicodedata

import altair as alt
import pandas as pd
import streamlit as st
from openpyxl import load_workbook


# Dashboard de avance para el anexo de construcción BESS.
# La hoja del archivo se llama exactamente "DUNE BESS CONSTR.".
NOMBRE_HOJA = "DUNE BESS CONSTR."
FILA_INICIO = 11
COL_PLAN_INICIO = 23   # W
COL_PLAN_FIN = 89      # CK
COL_REAL_INICIO = 92   # CN
COL_REAL_FIN = 158     # FB
TOLERANCIA = 0.0001
LOCALE_NUMERICO_ES = {
    "number": {
        "decimal": ",",
        "thousands": ".",
        "grouping": [3],
        "currency": ["", ""],
    }
}

st.set_page_config(
    page_title="Dashboard Construcción BESS",
    page_icon="📊",
    layout="wide",
)


def numero(valor, predeterminado=0.0):
    """Convierte celdas numéricas del Excel a float."""
    if valor is None:
        return predeterminado
    try:
        resultado = float(valor)
        return resultado if math.isfinite(resultado) else predeterminado
    except (TypeError, ValueError):
        return predeterminado


def separadores_es(valor_formateado):
    """Cambia el formato numérico de EE. UU. al formato español."""
    return valor_formateado.replace(",", "\0").replace(".", ",").replace("\0", ".")


def mostrar_numero(valor, decimales=0):
    """Muestra números enteros con punto como separador de miles."""
    return separadores_es(format(valor, f",.{decimales}f"))


def mostrar_porcentaje(valor, con_signo=False):
    """Muestra porcentajes con hasta dos decimales y separadores españoles."""
    porcentaje = valor * 100
    es_entero = math.isclose(porcentaje, round(porcentaje), rel_tol=0.0, abs_tol=1e-9)
    decimales = 0 if es_entero else 2
    signo = "+" if con_signo else ""
    formato = f"{signo},.{decimales}f"
    return f"{separadores_es(format(porcentaje, formato))}%"


def fecha_excel(valor):
    """Convierte fechas de Excel o fechas ya interpretadas por openpyxl."""
    if valor is None:
        return None
    if isinstance(valor, (datetime, date, pd.Timestamp)):
        return pd.Timestamp(valor).normalize()
    if isinstance(valor, (int, float)) and math.isfinite(float(valor)):
        return (pd.Timestamp("1899-12-30") + pd.to_timedelta(float(valor), unit="D")).normalize()
    fecha = pd.to_datetime(valor, dayfirst=True, errors="coerce")
    return None if pd.isna(fecha) else pd.Timestamp(fecha).normalize()


def es_si(valor):
    return str(valor or "").strip().casefold() in {"sí", "si", "yes", "true", "1"}


@st.cache_data(show_spinner="Leyendo la hoja de construcción...")
def cargar_datos(contenido_archivo):
    """Lee el .xlsm en modo de solo lectura y usa los resultados guardados en Excel."""
    libro = load_workbook(
        BytesIO(contenido_archivo),
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    if NOMBRE_HOJA not in libro.sheetnames:
        hojas_disponibles = ", ".join(libro.sheetnames)
        libro.close()
        raise ValueError(
            f'No se encontró la hoja "{NOMBRE_HOJA}". Hojas disponibles: {hojas_disponibles}'
        )

    hoja = libro[NOMBRE_HOJA]
    fecha_corte_archivo = fecha_excel(hoja["R8"].value)
    etiqueta_corte_archivo = hoja["S8"].value

    semanas = []
    for indice, columna in enumerate(range(COL_PLAN_INICIO, COL_PLAN_FIN + 1)):
        semanas.append(
            {
                "indice": indice,
                "etiqueta": str(hoja.cell(10, columna).value or f"Semana {indice + 1}"),
                "fecha": fecha_excel(hoja.cell(9, columna).value),
            }
        )

    filas = []
    ultima_fila = hoja.max_row or FILA_INICIO
    for numero_fila, celdas in enumerate(
        hoja.iter_rows(
            min_row=FILA_INICIO,
            max_row=ultima_fila,
            min_col=1,
            max_col=COL_REAL_FIN,
            values_only=True,
        ),
        start=FILA_INICIO,
    ):
        tarea = str(celdas[7] or "").strip()  # H: Tarea
        if not tarea:
            continue

        es_resumen = es_si(celdas[2])  # C: Resumen
        nivel = numero(celdas[1], predeterminado=-1)  # B: Nivel
        curva_plan = tuple(numero(celdas[col - 1]) for col in range(COL_PLAN_INICIO, COL_PLAN_FIN + 1))
        curva_real = tuple(numero(celdas[col - 1]) for col in range(COL_REAL_INICIO, COL_REAL_FIN + 1))

        filas.append(
            {
                "fila_excel": numero_fila,
                "nivel": nivel,
                "es_resumen": es_resumen,
                "id": celdas[5],  # F: Id
                "edt": str(celdas[6] or "").strip(),  # G: EDT
                "tarea": tarea,
                "inicio": fecha_excel(celdas[9]),  # J: Comienzo LB
                "fin": fecha_excel(celdas[10]),  # K: Fin LB
                "hh": numero(celdas[11]),  # L: Trabajo (h)
                "peso": numero(celdas[13]),  # N: Incidencia
                "curva_plan": curva_plan,
                "curva_real": curva_real,
            }
        )

    libro.close()
    if not semanas or not filas:
        raise ValueError("La hoja se abrió, pero no se encontraron semanas o actividades.")

    return {
        "semanas": semanas,
        "filas": filas,
        "fecha_corte_archivo": fecha_corte_archivo,
        "etiqueta_corte_archivo": etiqueta_corte_archivo,
    }


def etiqueta_alcance(fila):
    prefijo = f'{fila["edt"]} · ' if fila["edt"] else ""
    return f'{prefijo}{fila["tarea"]}'


def actividades_descendientes(filas, fila_resumen):
    """Devuelve las actividades hoja que pertenecen al resumen seleccionado."""
    posicion = next(
        (i for i, fila in enumerate(filas) if fila["fila_excel"] == fila_resumen["fila_excel"]),
        None,
    )
    if posicion is None:
        return []

    resultado = []
    for fila in filas[posicion + 1 :]:
        if fila["nivel"] >= 0 and fila["nivel"] <= fila_resumen["nivel"]:
            break
        if not fila["es_resumen"]:
            resultado.append(fila)
    return resultado



def normalizar_texto(texto):
    texto = unicodedata.normalize("NFKD", str(texto).casefold())
    return "".join(caracter for caracter in texto if not unicodedata.combining(caracter))


def contexto_actividades(filas):
    """Obtiene la ruta resumen, disciplina y Power Block de cada actividad hoja."""
    disciplinas = (
        ({"obras civiles", "obra civil"}, "Obras Civiles"),
        ({"obras mecanicas", "obra mecanica"}, "Obras Mec\u00e1nicas"),
        ({"obras electricas", "obra electrica"}, "Obras El\u00e9ctricas"),
    )
    patron_pb = re.compile(r"\bPB\s*[-#]?\s*(\d+)", re.IGNORECASE)
    pila_resumenes = []
    contextos = {}

    for fila in filas:
        nivel = fila["nivel"]
        while pila_resumenes and (nivel < 0 or pila_resumenes[-1]["nivel"] >= nivel):
            pila_resumenes.pop()

        if fila["es_resumen"]:
            pila_resumenes.append(fila)
            continue

        ruta = " \u203a ".join(etiqueta_alcance(resumen) for resumen in pila_resumenes)
        disciplina = None
        for resumen in reversed(pila_resumenes):
            nombre = normalizar_texto(resumen["tarea"])
            coincidencia = next(
                (etiqueta for alias, etiqueta in disciplinas if nombre in alias),
                None,
            )
            if coincidencia:
                disciplina = coincidencia
                break

        resumen_pb = next(
            (resumen for resumen in reversed(pila_resumenes) if patron_pb.search(resumen["tarea"])),
            None,
        )
        power_block = ""
        if resumen_pb:
            identificador = patron_pb.search(resumen_pb["tarea"]).group(1)
            power_block = f"PB-{identificador}"
            pb_fisico = re.search(
                r"\(\s*PB\s*[-#]?\s*(\d+)\s*\)",
                resumen_pb["tarea"],
                re.IGNORECASE,
            )
            if pb_fisico and pb_fisico.group(1) != identificador:
                power_block += f" (PB-{pb_fisico.group(1)})"

        contextos[fila["fila_excel"]] = {
            "ruta": ruta,
            "grupo": disciplina if resumen_pb and disciplina else ruta,
            "pb": power_block,
        }

    return contextos

st.title("Dashboard de Construcción BESS")
st.caption(
    'Fuente: hoja "DUNE BESS CONSTR." del archivo de avance. '
    "El libro se abre en modo de solo lectura."
)

archivo = st.sidebar.file_uploader(
    "Selecciona el archivo Excel de avance",
    type=["xlsm", "xlsx"],
    help="Carga el anexo que contiene la hoja DUNE BESS CONSTR.",
)

if archivo is None:
    st.info("Carga el archivo .xlsm o .xlsx para visualizar el dashboard.")
    st.stop()

try:
    datos = cargar_datos(archivo.getvalue())
except Exception as error:
    st.error(f"No fue posible leer el archivo: {error}")
    st.stop()

filas = datos["filas"]
semanas = datos["semanas"]
resumenes = [
    fila
    for fila in filas
    if fila["es_resumen"]
]
if not resumenes:
    st.error("No se encontraron filas resumen para construir el selector de alcance.")
    st.stop()

etiquetas = {fila["fila_excel"]: etiqueta_alcance(fila) for fila in resumenes}
ids_resumen = [fila["fila_excel"] for fila in resumenes]
resumen_por_id = {fila["fila_excel"]: fila for fila in resumenes}

with st.sidebar:
    st.header("Filtros")
    id_alcance = st.selectbox(
        "Actividad resumen (área / alcance)",
        options=ids_resumen,
        format_func=lambda valor: etiquetas[valor],
    )
    resumen_seleccionado = resumen_por_id[id_alcance]

    semanas_validas = [semana for semana in semanas if semana["fecha"] is not None]
    if not semanas_validas:
        st.error("No hay fechas semanales válidas en las filas 9 y 10 de la hoja.")
        st.stop()

    fecha_corte_archivo = datos["fecha_corte_archivo"]
    if fecha_corte_archivo is None:
        indice_defecto = len(semanas_validas) - 1
    else:
        indice_defecto = min(
            range(len(semanas_validas)),
            key=lambda i: abs((semanas_validas[i]["fecha"] - fecha_corte_archivo).days),
        )

    semana_seleccionada = st.selectbox(
        "Semana de control",
        options=range(len(semanas_validas)),
        index=indice_defecto,
        format_func=lambda i: (
            f'{semanas_validas[i]["etiqueta"]} · '
            f'{semanas_validas[i]["fecha"].strftime("%d-%m-%Y")}'
        ),
    )

semana = semanas_validas[semana_seleccionada]
indice_semana = semana["indice"]
fecha_corte = semana["fecha"]
avance_plan = max(0.0, min(1.0, resumen_seleccionado["curva_plan"][indice_semana]))
avance_real = max(0.0, min(1.0, resumen_seleccionado["curva_real"][indice_semana]))
desviacion = avance_real - avance_plan
spi = avance_real / avance_plan if avance_plan > TOLERANCIA else None

st.subheader(etiqueta_alcance(resumen_seleccionado))
st.caption(f"Semana de control: {semana['etiqueta']} · {fecha_corte:%d-%m-%Y}")

def tarjeta_indicador(titulo, valor, estado, detalle="", ayuda=""):
    colores = {
        "positivo": {"fondo": "#F0FDF4", "borde": "#BBF7D0", "texto": "#166534"},
        "negativo": {"fondo": "#FEF2F2", "borde": "#FECACA", "texto": "#B91C1C"},
        "neutro": {"fondo": "#F8FAFC", "borde": "#E2E8F0", "texto": "#334155"},
    }
    color = colores[estado]
    tarjeta = f"""
<div title="{escape(ayuda, quote=True)}" style="background:{color['fondo']}; border:1px solid {color['borde']}; border-radius:12px; padding:16px 18px; min-height:112px; box-shadow:0 2px 7px rgba(15, 23, 42, 0.05);">
  <div style="font-size:0.84rem; font-weight:600; color:#475569; margin-bottom:10px;">{escape(titulo)}</div>
  <div style="font-size:1.85rem; line-height:1.15; font-weight:750; color:{color['texto']};">{escape(str(valor))}</div>
  <div style="font-size:0.76rem; color:#64748B; margin-top:8px;">{escape(detalle)}</div>
</div>
"""
    st.markdown(tarjeta, unsafe_allow_html=True)

st.markdown("##### Indicadores a la fecha de control")
col_plan, col_real, col_desviacion, col_spi = st.columns(4, gap="medium")

estado_plan = "positivo" if avance_plan > TOLERANCIA else "neutro"
estado_real = "positivo" if avance_real > TOLERANCIA else "neutro"
if desviacion > TOLERANCIA:
    estado_desviacion = "positivo"
    detalle_desviacion = "Adelanto respecto del plan"
elif desviacion < -TOLERANCIA:
    estado_desviacion = "negativo"
    detalle_desviacion = "Atraso respecto del plan"
else:
    estado_desviacion = "neutro"
    detalle_desviacion = "Sin desviaci\u00f3n significativa"

if spi is None:
    estado_spi = "neutro"
    detalle_spi = "No disponible: avance planificado igual a 0"
elif spi < 0.80:
    estado_spi = "negativo"
    detalle_spi = "Bajo el umbral de alerta (0,80)"
else:
    estado_spi = "positivo"
    detalle_spi = "Umbral de alerta: 0,80"

with col_plan:
    tarjeta_indicador("% Planificado", mostrar_porcentaje(avance_plan), estado_plan, "Meta acumulada")
with col_real:
    tarjeta_indicador("% Avance real", mostrar_porcentaje(avance_real), estado_real, "Avance ejecutado acumulado")
with col_desviacion:
    tarjeta_indicador(
        "% Desviaci\u00f3n",
        mostrar_porcentaje(desviacion, con_signo=True),
        estado_desviacion,
        detalle_desviacion,
    )
with col_spi:
    tarjeta_indicador(
        "SPI",
        mostrar_numero(spi, 2) if spi is not None else "N/D",
        estado_spi,
        detalle_spi,
        "SPI = avance real / avance planificado. SPI menor a 0,80 se muestra en rojo.",
    )

tab_dashboard, tab_actividades, tab_atrasadas = st.tabs(
    ["1. Dashboard y curva S", "2. Estado de actividades", "3. Atrasadas por HH"]
)

with tab_dashboard:
    st.markdown("#### Curva S acumulada")
    curva = pd.DataFrame(
        {
            "Fecha": [s["fecha"] for s in semanas_validas],
            "Semana": [s["etiqueta"] for s in semanas_validas],
            "Planificado": [
                max(0.0, min(1.0, resumen_seleccionado["curva_plan"][s["indice"]]))
                for s in semanas_validas
            ],
            "Avance real": [
                (
                    max(0.0, min(1.0, resumen_seleccionado["curva_real"][s["indice"]]))
                    if s["indice"] <= indice_semana
                    else None
                )
                for s in semanas_validas
            ],
        }
    )
    curva = curva.sort_values("Fecha")
    curva_larga = curva.melt(
        id_vars=["Fecha", "Semana"],
        var_name="Serie",
        value_name="Avance",
    )
    grafico_curva = (
        alt.Chart(curva_larga)
        .mark_line()
        .encode(
            x=alt.X(
                "Fecha:T",
                title="Fecha",
                axis=alt.Axis(format="%d-%m-%Y", labelAngle=-90),
            ),
            y=alt.Y(
                "Avance:Q",
                title="Avance acumulado",
                scale=alt.Scale(domain=[0, 1]),
                axis=alt.Axis(format=".2~%"),
            ),
            color=alt.Color("Serie:N", title="Curva"),
            tooltip=[
                alt.Tooltip("Semana:N", title="Semana"),
                alt.Tooltip("Fecha:T", title="Fecha", format="%d-%m-%Y"),
                alt.Tooltip("Serie:N", title="Curva"),
                alt.Tooltip("Avance:Q", title="Avance", format=".2~%"),
            ],
        )
    )
    etiqueta_data_date = f"DATA DATE · {fecha_corte:%d-%m-%Y}"
    alineacion_etiqueta = (
        "left"
        if (fecha_corte - curva["Fecha"].min()).days
        <= (curva["Fecha"].max() - fecha_corte).days
        else "right"
    )
    desplazamiento_etiqueta = 8 if alineacion_etiqueta == "left" else -8
    datos_data_date = pd.DataFrame({"Fecha": [fecha_corte]})
    regla_data_date = (
        alt.Chart(datos_data_date)
        .mark_rule(color="#B45309", strokeDash=[6, 4], strokeWidth=2)
        .encode(x=alt.X("Fecha:T"))
    )
    datos_etiqueta_data_date = pd.DataFrame(
        {"Fecha": [fecha_corte], "Nivel": [0.96], "Etiqueta": [etiqueta_data_date]}
    )
    texto_data_date = (
        alt.Chart(datos_etiqueta_data_date)
        .mark_text(
            align=alineacion_etiqueta,
            baseline="top",
            dx=desplazamiento_etiqueta,
            dy=8,
            color="#7C2D12",
            fontSize=13,
            fontWeight="bold",
        )
        .encode(
            x=alt.X("Fecha:T"),
            y=alt.Y("Nivel:Q", scale=alt.Scale(domain=[0, 1]), axis=None),
            text=alt.Text("Etiqueta:N"),
        )
    )
    grafico = (
        alt.layer(grafico_curva, regla_data_date, texto_data_date)
        .resolve_scale(y="shared")
        .configure(locale=LOCALE_NUMERICO_ES)
    )
    st.altair_chart(
        grafico,
        width="stretch",
        height=380,
        alt="Curva S acumulada planificada y real con línea y etiqueta DATA DATE en la semana de control.",
    )
    st.caption(
        "La selección de alcance y semana de control actualiza los indicadores y la tabla. "
        "La línea DATA DATE señala la fecha de control seleccionada. "
        "La desviación se calcula como avance real menos avance planificado."
    )

    st.markdown("#### Avance parcial semanal: últimas 6 semanas")
    semanas_hasta_corte = sorted(
        [s for s in semanas_validas if s["indice"] <= indice_semana],
        key=lambda s: s["indice"],
    )
    ventana_6_semanas = semanas_hasta_corte[-6:]
    registros_parciales = []
    for semana_parcial in ventana_6_semanas:
        indice = semana_parcial["indice"]
        indice_anterior = indice - 1
        plan_actual = resumen_seleccionado["curva_plan"][indice]
        real_actual = resumen_seleccionado["curva_real"][indice]
        plan_anterior = (
            resumen_seleccionado["curva_plan"][indice_anterior]
            if indice_anterior >= 0
            else 0.0
        )
        real_anterior = (
            resumen_seleccionado["curva_real"][indice_anterior]
            if indice_anterior >= 0
            else 0.0
        )
        avance_plan_parcial = plan_actual - plan_anterior
        avance_real_parcial = real_actual - real_anterior
        registros_parciales.append(
            {
                "Semana": semana_parcial["etiqueta"],
                "Fecha": semana_parcial["fecha"],
                "Planificado": avance_plan_parcial,
                "Avance real": avance_real_parcial,
                "Desviacion": avance_real_parcial - avance_plan_parcial,
            }
        )

    avance_parcial_ancho = pd.DataFrame(registros_parciales)
    avance_parcial = avance_parcial_ancho.melt(
        id_vars=["Semana", "Fecha"],
        value_vars=["Planificado", "Avance real"],
        var_name="Serie",
        value_name="Avance_parcial",
    )
    avance_parcial["Etiqueta"] = avance_parcial["Avance_parcial"].map(mostrar_porcentaje)

    grafico_lineas = (
        alt.Chart(avance_parcial)
        .mark_line(point=True)
        .encode(
            x=alt.X(
                "Fecha:T",
                title="Fecha",
                axis=alt.Axis(format="%d-%m-%Y", labelAngle=-90),
            ),
            y=alt.Y(
                "Avance_parcial:Q",
                title="Avance parcial / desviación semanal",
                axis=alt.Axis(format=".2~%"),
            ),
            color=alt.Color(
                "Serie:N",
                title="Curva",
                scale=alt.Scale(
                    domain=["Planificado", "Avance real"],
                    range=["#1D4ED8", "#C2410C"],
                ),
            ),
            detail=alt.Detail("Serie:N"),
            tooltip=[
                alt.Tooltip("Semana:N", title="Semana"),
                alt.Tooltip("Fecha:T", title="Fecha", format="%d-%m-%Y"),
                alt.Tooltip("Serie:N", title="Curva"),
                alt.Tooltip("Avance_parcial:Q", title="Avance parcial", format=".2~%"),
            ],
        )
    )
    etiqueta_plan = (
        alt.Chart(avance_parcial)
        .transform_filter(alt.datum.Serie == "Planificado")
        .mark_text(
            dy=-13,
            color="#111827",
            fontSize=12,
            fontWeight="bold",
        )
        .encode(
            x=alt.X("Fecha:T"),
            y=alt.Y("Avance_parcial:Q"),
            text=alt.Text("Etiqueta:N"),
        )
    )
    etiqueta_real = (
        alt.Chart(avance_parcial)
        .transform_filter(alt.datum.Serie == "Avance real")
        .mark_text(
            dy=13,
            color="#111827",
            fontSize=12,
            fontWeight="bold",
        )
        .encode(
            x=alt.X("Fecha:T"),
            y=alt.Y("Avance_parcial:Q"),
            text=alt.Text("Etiqueta:N"),
        )
    )
    desviacion_semanal = avance_parcial_ancho.copy()
    desviacion_semanal["Resultado"] = desviacion_semanal["Desviacion"].map(
        lambda valor: (
            "Adelantada"
            if valor > TOLERANCIA
            else "Atrasada"
            if valor < -TOLERANCIA
            else "En línea"
        )
    )
    desviacion_semanal["Etiqueta"] = desviacion_semanal["Desviacion"].map(
        lambda valor: (
            "0%"
            if math.isclose(valor, 0.0, rel_tol=0.0, abs_tol=TOLERANCIA)
            else mostrar_porcentaje(valor, con_signo=True)
        )
    )
    margen_etiqueta = max(
        desviacion_semanal["Desviacion"].abs().max() * 0.1,
        0.002,
    )
    desviacion_semanal["Nivel_etiqueta"] = desviacion_semanal["Desviacion"].map(
        lambda valor: valor + margen_etiqueta if valor >= 0 else valor - margen_etiqueta
    )
    escala_estado = alt.Scale(
        domain=["Adelantada", "Atrasada", "En línea"],
        range=["#15803D", "#B91C1C", "#64748B"],
    )
    eje_fecha_desviacion = alt.X(
        "Fecha:T",
        title="Fecha",
        axis=alt.Axis(format="%d-%m-%Y", labelAngle=-90),
    )
    barras_desviacion = (
        alt.Chart(desviacion_semanal)
        .mark_bar(size=24, opacity=0.32)
        .encode(
            x=eje_fecha_desviacion,
            y=alt.Y("Desviacion:Q", axis=None),
            color=alt.Color("Resultado:N", title="Estado", scale=escala_estado),
            tooltip=[
                alt.Tooltip("Semana:N", title="Semana"),
                alt.Tooltip("Fecha:T", title="Fecha", format="%d-%m-%Y"),
                alt.Tooltip("Resultado:N", title="Estado"),
                alt.Tooltip("Etiqueta:N", title="Desviación"),
            ],
        )
    )
    linea_cero = (
        alt.Chart(pd.DataFrame({"Cero": [0.0]}))
        .mark_rule(color="#AAB2BF", strokeDash=[3, 3], strokeWidth=1)
        .encode(y=alt.Y("Cero:Q", axis=None))
    )
    etiquetas_desviacion = (
        alt.Chart(desviacion_semanal)
        .mark_text(fontSize=11, fontWeight="bold")
        .encode(
            x=alt.X("Fecha:T"),
            y=alt.Y("Nivel_etiqueta:Q", axis=None),
            text=alt.Text("Etiqueta:N"),
            color=alt.Color("Resultado:N", scale=escala_estado, legend=None),
        )
    )
    grafico_parcial = (
        alt.layer(
            barras_desviacion,
            linea_cero,
            grafico_lineas,
            etiqueta_plan,
            etiqueta_real,
            etiquetas_desviacion,
        )
        .resolve_scale(y="shared")
        .configure(locale=LOCALE_NUMERICO_ES)
    )
    st.altair_chart(
        grafico_parcial,
        width="stretch",
        height=410,
        alt="Gráfico integrado de líneas para avance parcial planificado y real, con barras de desviación semanal. Verde indica adelanto y rojo atraso; las series y desviaciones tienen etiquetas de valor.",
    )
    st.caption(
        "Cada punto muestra el incremento de avance respecto de la semana anterior. "
        "Las barras comparan el avance real semanal con el planificado. "
        f'La última semana del gráfico es la semana de control seleccionada: {semana["etiqueta"]}.'
    )

with tab_actividades:
    st.markdown("#### Estado de actividades por disciplina")
    st.caption("Los Power Blocks se agrupan por disciplina y se ordenan por n\u00famero de PB.")
    actividades = actividades_descendientes(filas, resumen_seleccionado)
    contextos = contexto_actividades(filas)
    registros = []

    for actividad in actividades:
        plan = max(0.0, min(1.0, actividad["curva_plan"][indice_semana]))
        real = max(0.0, min(1.0, actividad["curva_real"][indice_semana]))

        if real <= TOLERANCIA and actividad["inicio"] is not None and fecha_corte < actividad["inicio"]:
            estado = "Sin iniciar"
        elif real + TOLERANCIA < plan:
            estado = "Atrasada"
        else:
            estado = "No atrasada"

        contexto = contextos.get(actividad["fila_excel"], {})
        registros.append(
            {
                "Grupo": contexto.get("grupo") or etiqueta_alcance(resumen_seleccionado),
                "PB": contexto.get("pb", ""),
                "Actividad resumen": contexto.get("ruta") or etiqueta_alcance(resumen_seleccionado),
                "ID": actividad["id"],
                "EDT": actividad["edt"],
                "Actividad": actividad["tarea"],
                "Inicio LB": actividad["inicio"].strftime("%d-%m-%Y") if actividad["inicio"] is not None else "",
                "Fin LB": actividad["fin"].strftime("%d-%m-%Y") if actividad["fin"] is not None else "",
                "HH": actividad["hh"],
                "% Planificado": plan,
                "% Avance real": real,
                "Desviación (%)": real - plan,
                "Estado": estado,
            }
        )

    tabla = pd.DataFrame(registros)
    if tabla.empty:
        st.info("Este alcance no tiene actividades hoja para mostrar.")
    else:
        conteo = Counter(tabla["Estado"])
        col_atrasadas, col_no_atrasadas, col_sin_iniciar = st.columns(3)
        col_atrasadas.metric("Atrasadas", mostrar_numero(conteo.get("Atrasada", 0), 0))
        col_no_atrasadas.metric("No atrasadas", mostrar_numero(conteo.get("No atrasada", 0), 0))
        col_sin_iniciar.metric("Sin iniciar", mostrar_numero(conteo.get("Sin iniciar", 0), 0))

        filtro_estado = st.selectbox(
            "Filtrar actividades por estado",
            ["Todas", "Atrasada", "No atrasada", "Sin iniciar"],
        )
        busqueda = st.text_input("Buscar actividad, resumen, EDT o ID")

        tabla_visible = tabla.copy()
        if filtro_estado != "Todas":
            tabla_visible = tabla_visible[tabla_visible["Estado"] == filtro_estado]
        if busqueda.strip():
            texto = busqueda.strip()
            mascara = (
                tabla_visible["Actividad"].astype(str).str.contains(texto, case=False, na=False)
                | tabla_visible["Actividad resumen"].astype(str).str.contains(texto, case=False, na=False)
                | tabla_visible["EDT"].astype(str).str.contains(texto, case=False, na=False)
                | tabla_visible["ID"].astype(str).str.contains(texto, case=False, na=False)
            )
            tabla_visible = tabla_visible[mascara]

        columnas_porcentaje = ["% Planificado", "% Avance real", "Desviaci\u00f3n (%)"]
        orden_disciplina = {
            "Obras Civiles": 0,
            "Obras Mec\u00e1nicas": 1,
            "Obras El\u00e9ctricas": 2,
        }
        grupos = list(tabla_visible.groupby("Grupo", sort=False))
        grupos.sort(key=lambda elemento: orden_disciplina.get(elemento[0], 3))

        for nombre_grupo, grupo in grupos:
            es_disciplina = nombre_grupo in orden_disciplina
            if es_disciplina:
                titulo = nombre_grupo
                ruta_padre = ""
            else:
                partes_ruta = nombre_grupo.split(" \u203a ")
                titulo = partes_ruta[-1]
                ruta_padre = " \u203a ".join(partes_ruta[:-1])

            columnas_ocultas = ["Actividad resumen", "Grupo"]
            if not es_disciplina:
                columnas_ocultas.append("PB")
            grupo_visible = grupo.drop(columns=columnas_ocultas).copy()
            if es_disciplina:
                grupo_visible["_orden_pb"] = grupo_visible["PB"].map(
                    lambda valor: (
                        int(coincidencia.group(1))
                        if (coincidencia := re.search(r"PB-(\d+)", str(valor), re.IGNORECASE))
                        else math.inf
                    )
                )
                grupo_visible = grupo_visible.sort_values(
                    ["_orden_pb", "Actividad"],
                    kind="stable",
                ).drop(columns=["_orden_pb"])

            grupo_formateado = grupo_visible.style.format(
                {
                    "HH": mostrar_numero,
                    **{columna: mostrar_porcentaje for columna in columnas_porcentaje},
                }
            )
            with st.container(border=True):
                st.markdown(
                    f'<div style="font-size:1.05rem;font-weight:700;color:#0F172A;">'
                    f'{escape(titulo)}</div>',
                    unsafe_allow_html=True,
                )
                if ruta_padre:
                    st.markdown(
                        f'<div style="font-size:0.8rem;color:#64748B;margin-top:3px;">'
                        f'Ruta: {escape(ruta_padre)}</div>',
                        unsafe_allow_html=True,
                    )
                st.caption(
                    f"{len(grupo_visible)} actividades  |  "
                    f"{mostrar_numero(grupo_visible['HH'].sum())} HH"
                )
                st.dataframe(
                    grupo_formateado,
                    hide_index=True,
                    width="stretch",
                    alt="Actividades agrupadas por disciplina y Power Block.",
                )
    st.caption(
        "Estado de atraso: "
        "Sin iniciar cuando el avance real es cero y la fecha de control es anterior al inicio LB; "
        "Atrasada cuando el avance real está por debajo del plan; No atrasada en los demás casos."
    )

with tab_atrasadas:
    st.markdown("#### Actividades atrasadas por disciplina y HH")
    st.caption("Los Power Blocks se consolidan por disciplina y se identifican en la columna PB.")
    if tabla.empty:
        st.info("Este alcance no tiene actividades para revisar.")
    else:
        atrasadas = (
            tabla.loc[tabla["Estado"] == "Atrasada"]
            .sort_values("HH", ascending=False)
            .reset_index(drop=True)
        )
        if atrasadas.empty:
            st.success("No hay actividades atrasadas para el alcance y semana seleccionados.")
        else:
            total_hh_atrasadas = atrasadas["HH"].sum()
            actividad_mayor_hh = atrasadas.iloc[0]
            col_cantidad, col_hh, col_mayor = st.columns(3)
            col_cantidad.metric("Actividades atrasadas", mostrar_numero(len(atrasadas), 0))
            col_hh.metric("HH atrasadas", mostrar_numero(total_hh_atrasadas))
            col_mayor.metric("Mayor HH", mostrar_numero(actividad_mayor_hh["HH"]))
            referencia_pb = f' [{actividad_mayor_hh["PB"]}]' if actividad_mayor_hh["PB"] else ""
            st.caption(
                f'Actividad con mayor HH: **{actividad_mayor_hh["Actividad"]}**{referencia_pb} '
                f'({mostrar_numero(actividad_mayor_hh["HH"])} h). '
                "Las tablas se ordenan de mayor a menor segun Trabajo (h) del Excel."
            )

            columnas_porcentaje = ["% Planificado", "% Avance real", "Desviaci\u00f3n (%)"]
            orden_disciplina = {
                "Obras Civiles": 0,
                "Obras Mec\u00e1nicas": 1,
                "Obras El\u00e9ctricas": 2,
            }
            grupos = list(atrasadas.groupby("Grupo", sort=False))
            grupos.sort(key=lambda elemento: orden_disciplina.get(elemento[0], 3))

            for nombre_grupo, grupo in grupos:
                es_disciplina = nombre_grupo in orden_disciplina
                if es_disciplina:
                    titulo = nombre_grupo
                    ruta_padre = ""
                else:
                    partes_ruta = nombre_grupo.split(" \u203a ")
                    titulo = partes_ruta[-1]
                    ruta_padre = " \u203a ".join(partes_ruta[:-1])

                columnas_ocultas = ["Actividad resumen", "Grupo"]
                if not es_disciplina:
                    columnas_ocultas.append("PB")
                grupo_visible = grupo.drop(columns=columnas_ocultas).copy()
                grupo_formateado = grupo_visible.style.format(
                    {
                        "HH": mostrar_numero,
                        **{columna: mostrar_porcentaje for columna in columnas_porcentaje},
                    }
                )

                with st.container(border=True):
                    st.markdown(
                        f'<div style="font-size:1.05rem;font-weight:700;color:#0F172A;">'
                        f'{escape(titulo)}</div>',
                        unsafe_allow_html=True,
                    )
                    if ruta_padre:
                        st.markdown(
                            f'<div style="font-size:0.8rem;color:#64748B;margin-top:3px;">'
                            f'Ruta: {escape(ruta_padre)}</div>',
                            unsafe_allow_html=True,
                        )
                    st.caption(
                        f"{len(grupo_visible)} actividades atrasadas  |  "
                        f"{mostrar_numero(grupo_visible['HH'].sum())} HH"
                    )
                    st.dataframe(
                        grupo_formateado,
                        hide_index=True,
                        width="stretch",
                        alt="Actividades atrasadas ordenadas por HH y agrupadas por disciplina.",
                    )

